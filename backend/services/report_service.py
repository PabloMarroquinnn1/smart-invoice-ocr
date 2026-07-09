import os
import csv
from datetime import datetime
from fpdf import FPDF
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

REPORTS_DIR = os.path.join(os.path.dirname(os.path.dirname(__file__)), "reports")
os.makedirs(REPORTS_DIR, exist_ok=True)


def generar_reporte_pdf(facturas: list) -> str:
    pdf = FPDF()
    pdf.add_page()
    pdf.set_auto_page_break(auto=True, margin=15)

    pdf.set_font("Helvetica", "B", 18)
    pdf.cell(0, 12, "SmartInvoice - Reporte de Facturas", ln=True, align="C")
    pdf.set_font("Helvetica", "", 10)
    pdf.cell(0, 8, f"Generado: {datetime.now().strftime('%d/%m/%Y %H:%M')}", ln=True, align="C")
    pdf.ln(10)

    total_facturas = len(facturas)
    monto_total = sum(f.total for f in facturas)
    total_iva = sum(f.impuestos for f in facturas)

    pdf.set_font("Helvetica", "B", 12)
    pdf.cell(0, 8, "Resumen General", ln=True)
    pdf.set_font("Helvetica", "", 10)
    pdf.cell(0, 6, f"Total de facturas: {total_facturas}", ln=True)
    pdf.cell(0, 6, f"Monto total: Q {monto_total:,.2f}", ln=True)
    pdf.cell(0, 6, f"IVA total: Q {total_iva:,.2f}", ln=True)
    pdf.ln(8)

    pdf.set_font("Helvetica", "B", 9)
    col_widths = [25, 45, 25, 22, 25, 25, 25]
    headers = ["No. Factura", "Proveedor", "NIT", "Fecha", "Subtotal", "IVA", "Total"]

    pdf.set_fill_color(48, 43, 99)
    pdf.set_text_color(255, 255, 255)
    for i, h in enumerate(headers):
        pdf.cell(col_widths[i], 8, h, border=1, fill=True, align="C")
    pdf.ln()

    pdf.set_font("Helvetica", "", 8)
    pdf.set_text_color(0, 0, 0)

    for f in facturas:
        pdf.cell(col_widths[0], 7, str(f.numero_factura)[:12], border=1, align="C")
        pdf.cell(col_widths[1], 7, str(f.nombre_proveedor or '-')[:22], border=1)
        pdf.cell(col_widths[2], 7, str(f.nit_proveedor or '-')[:12], border=1, align="C")
        pdf.cell(col_widths[3], 7, str(f.fecha or '-')[:10], border=1, align="C")
        pdf.cell(col_widths[4], 7, f"Q {f.subtotal:,.2f}", border=1, align="R")
        pdf.cell(col_widths[5], 7, f"Q {f.impuestos:,.2f}", border=1, align="R")
        pdf.cell(col_widths[6], 7, f"Q {f.total:,.2f}", border=1, align="R")
        pdf.ln()

    filename = f"reporte_{datetime.now().strftime('%Y%m%d_%H%M%S')}.pdf"
    filepath = os.path.join(REPORTS_DIR, filename)
    pdf.output(filepath)
    return filepath


def generar_reporte_excel(facturas: list) -> str:
    wb = Workbook()
    ws = wb.active
    ws.title = "Facturas"

    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="302B63", end_color="302B63", fill_type="solid")
    thin_border = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin")
    )

    headers = ["No. Factura", "Proveedor", "NIT", "Fecha", "Cliente", "Subtotal", "IVA 12%", "Total", "Estado"]
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
        cell.border = thin_border

    for row, f in enumerate(facturas, 2):
        data = [
            f.numero_factura, f.nombre_proveedor or '-', f.nit_proveedor or '-',
            f.fecha or '-', f.cliente_nombre or '-',
            f.subtotal, f.impuestos, f.total, f.estado
        ]
        for col, val in enumerate(data, 1):
            cell = ws.cell(row=row, column=col, value=val)
            cell.border = thin_border
            if col in [6, 7, 8]:
                cell.number_format = '#,##0.00'

    widths = [15, 35, 15, 12, 30, 15, 15, 15, 12]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + i)].width = w

    last_row = len(facturas) + 2
    ws.cell(row=last_row, column=5, value="TOTALES:").font = Font(bold=True)
    for col, formula_col in [(6, 'F'), (7, 'G'), (8, 'H')]:
        cell = ws.cell(row=last_row, column=col)
        cell.value = f"=SUM({formula_col}2:{formula_col}{last_row - 1})"
        cell.font = Font(bold=True)
        cell.number_format = '#,##0.00'

    filename = f"reporte_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    filepath = os.path.join(REPORTS_DIR, filename)
    wb.save(filepath)
    return filepath


def generar_reporte_csv(facturas: list) -> str:
    filename = f"reporte_{datetime.now().strftime('%Y%m%d_%H%M%S')}.csv"
    filepath = os.path.join(REPORTS_DIR, filename)

    with open(filepath, 'w', newline='', encoding='utf-8') as f:
        writer = csv.writer(f)
        writer.writerow(["No. Factura", "Proveedor", "NIT", "Fecha", "Cliente", "Subtotal", "IVA", "Total", "Estado"])
        for fac in facturas:
            writer.writerow([
                fac.numero_factura, fac.nombre_proveedor or '', fac.nit_proveedor or '',
                fac.fecha or '', fac.cliente_nombre or '',
                fac.subtotal, fac.impuestos, fac.total, fac.estado
            ])

    return filepath
